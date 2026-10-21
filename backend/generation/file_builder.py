"""Builds real downloadable files (xlsx / csv / docx / pdf) from a JSON spec
that the LLM emits inside a fenced ```file ... ``` block. Fully offline."""
import csv
import json
import os
import re
import uuid

from core.config import settings

GENERATED_DIR = os.path.join(os.path.dirname(settings.UPLOAD_DIR.rstrip("/")), "generated")
FILE_BLOCK_RE = re.compile(r"```file\s*\n(.*?)\n?```", re.DOTALL)

_SAFE = re.compile(r"[^A-Za-z0-9._-]+")


def extract_file_directive(text):
    """Return (clean_text, spec|None). The directive block is removed from text."""
    m = FILE_BLOCK_RE.search(text or "")
    if not m:
        return text, None
    clean = (text[:m.start()] + text[m.end():]).strip()
    try:
        spec = json.loads(m.group(1))
        if not isinstance(spec, dict):
            return clean, None
    except Exception:
        return clean, None
    return clean, spec


def _base(spec, ext):
    os.makedirs(GENERATED_DIR, exist_ok=True)
    name = _SAFE.sub("_", os.path.splitext(str(spec.get("filename") or spec.get("title") or "document"))[0]).strip("._") or "document"
    return os.path.join(GENERATED_DIR, f"{uuid.uuid4().hex[:10]}_{name}.{ext}")


def _cols_rows(spec):
    cols = [str(c) for c in (spec.get("columns") or [])]
    rows = [["" if v is None else v for v in r] for r in (spec.get("rows") or []) if isinstance(r, (list, tuple))]
    if not cols and rows:
        cols = [f"Column {i+1}" for i in range(len(rows[0]))]
    return cols, rows


def build_csv(spec):
    path = _base(spec, "csv")
    cols, rows = _cols_rows(spec)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        if cols:
            w.writerow(cols)
        w.writerows(rows)
    return path


def build_xlsx(spec):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    path = _base(spec, "xlsx")
    cols, rows = _cols_rows(spec)
    wb = Workbook()
    ws = wb.active
    ws.title = (str(spec.get("title") or "Sheet1"))[:31].replace("/", "-")
    thin = Side(style="thin", color="BBBBBB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    if cols:
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="002060")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = border
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = border
            if isinstance(c.value, str) and c.value.startswith("="):
                pass  # keep formulas live
    for i in range(1, ws.max_column + 1):
        width = max((len(str(c.value)) if c.value is not None else 0) for c in ws[get_column_letter(i)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 3, 10), 60)
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


def build_docx(spec):
    from docx import Document
    path = _base(spec, "docx")
    d = Document()
    if spec.get("title"):
        d.add_heading(str(spec["title"]), 0)
    for s in spec.get("sections") or []:
        if not isinstance(s, dict):
            continue
        if s.get("heading"):
            d.add_heading(str(s["heading"]), 1)
        for para in str(s.get("body") or "").split("\n"):
            if para.strip():
                d.add_paragraph(para.strip())
    cols, rows = _cols_rows(spec)
    if cols and rows:
        t = d.add_table(rows=1, cols=len(cols))
        t.style = "Light Grid Accent 1"
        for i, c in enumerate(cols):
            t.rows[0].cells[i].text = c
        for r in rows:
            cells = t.add_row().cells
            for i in range(len(cols)):
                cells[i].text = str(r[i]) if i < len(r) else ""
    d.save(path)
    return path


def build_pdf(spec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape
    path = _base(spec, "pdf")
    ss = getSampleStyleSheet()
    story = []
    if spec.get("title"):
        story += [Paragraph(escape(str(spec["title"])), ss["Title"]), Spacer(1, 12)]
    for s in spec.get("sections") or []:
        if not isinstance(s, dict):
            continue
        if s.get("heading"):
            story.append(Paragraph(escape(str(s["heading"])), ss["Heading2"]))
        for para in str(s.get("body") or "").split("\n"):
            if para.strip():
                story += [Paragraph(escape(para.strip()), ss["BodyText"]), Spacer(1, 4)]
    cols, rows = _cols_rows(spec)
    if cols and rows:
        data = [[Paragraph(escape(str(c)), ss["BodyText"]) for c in cols]] + \
               [[Paragraph(escape(str(v)), ss["BodyText"]) for v in r] for r in rows]
        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#CDDEBF")),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story += [Spacer(1, 8), t]
    if not story:
        story.append(Paragraph("(empty document)", ss["BodyText"]))
    SimpleDocTemplate(path, pagesize=A4).build(story)
    return path


BUILDERS = {
    "xlsx": ("xlsx", build_xlsx), "excel": ("xlsx", build_xlsx), "spreadsheet": ("xlsx", build_xlsx),
    "csv": ("csv", build_csv),
    "docx": ("docx", build_docx), "word": ("docx", build_docx), "doc": ("docx", build_docx), "document": ("docx", build_docx),
    "pdf": ("pdf", build_pdf), "report": ("pdf", build_pdf),
}


def build_file(spec):
    entry = BUILDERS.get(str(spec.get("type", "")).lower().strip())
    if not entry:
        return None, None
    ftype, fn = entry
    return fn(spec), ftype
